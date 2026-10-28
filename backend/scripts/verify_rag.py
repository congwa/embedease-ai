"""Isolated direct functional acceptance; no production DB or bucket is used."""

import asyncio
import hashlib
import io
import secrets
import tempfile
from pathlib import Path

import httpx
from docx import Document as WordDocument
from fastapi import FastAPI, Request
from fastapi.responses import JSONResponse
from loguru import logger
from pypdf import PdfWriter
from sqlalchemy import select, update

from app.core.config import settings
from app.core.database import get_db_context
from app.core.db import provider
from app.models import Base
from app.models.agent import Agent, KnowledgeConfig
from app.models.conversation import Conversation
from app.models.knowledge_document import Document, DocumentChunk, DocumentRevision, IngestionJob
from app.models.message import Message
from app.models.user import User
from app.routers.agents import knowledge_router
from app.routers.knowledge_documents import citation_router, router
from app.services.knowledge.rag import service, worker
from app.services.knowledge.rag.citations import history_metadata
from app.services.knowledge.rag.settings import RagError, uid
from app.services.knowledge.rag.storage import DocumentStorage
from app.services.knowledge.rag.vectors import qdrant, scope_filter

logger.remove()


def pdf_bytes():
    """生成可提取文字的最小 PDF，避免依赖额外制图软件。"""
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(600, 800)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 750 Td (PDF acceptance: warranty period is 19 months.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


async def main():
    """临时 SQL、独立私有桶和随机知识源下执行真实服务验收并清理。"""
    tmp = tempfile.TemporaryDirectory(prefix="embedease-rag-")
    settings.KNOWLEDGE_ADMIN_KEY = secrets.token_hex(32)
    settings.KNOWLEDGE_BUCKET = "rag-qa-" + secrets.token_hex(5)
    settings.CHECKPOINT_DB_PATH = str(Path(tmp.name) / "checkpoints.db")
    # 验收只覆盖 RAG，关闭跨轮长期记忆并隔离其路径，避免后台抽取触及真实数据。
    settings.MEMORY_ENABLED = False
    settings.MEMORY_ORCHESTRATION_ENABLED = False
    settings.MEMORY_STORE_DB_PATH = str(Path(tmp.name) / "memory.db")
    settings.MEMORY_FACT_DB_PATH = str(Path(tmp.name) / "facts.db")
    settings.MEMORY_GRAPH_FILE_PATH = str(Path(tmp.name) / "graph.jsonl")
    settings.DATABASE_BACKEND = "sqlite"
    provider._provider = provider.SQLiteProvider(
        "sqlite+aiosqlite:///" + str(Path(tmp.name) / "app.db")
    )
    await provider._provider.init_db(Base)
    config_id, other_id, agent_id = uid(), uid(), uid()
    collections = set()
    storage = DocumentStorage()
    app = FastAPI()
    app.include_router(router)
    app.include_router(citation_router)
    app.include_router(knowledge_router)

    @app.exception_handler(RagError)
    async def rag_error(req: Request, exc: RagError):
        """与正式应用保持相同的业务错误状态。"""
        return JSONResponse({"detail": str(exc)}, status_code=exc.status)

    async with get_db_context() as db:
        db.add_all(
            [
                KnowledgeConfig(
                    id=config_id,
                    name="isolated QA",
                    type="vector",
                    collection_name="rag_qa_empty_" + uid(),
                ),
                KnowledgeConfig(
                    id=other_id,
                    name="isolation",
                    type="vector",
                    collection_name="rag_qa_empty_" + uid(),
                ),
            ]
        )
    base = f"/api/v1/admin/knowledge-configs/{config_id}"
    headers = {"X-Knowledge-Admin-Key": settings.KNOWLEDGE_ADMIN_KEY}
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://qa", headers=headers, timeout=120
    ) as client:
        try:
            res = await client.get(base + "/documents", headers={"X-Knowledge-Admin-Key": "wrong"})
            assert res.status_code == 401, res.text
            res = await client.post(base + "/rag/enable")
            assert res.status_code == 200, res.text
            await service.enable(other_id)
            print("ACCESS_AND_ENABLE_OK", flush=True)
            txt = (
                "入库验收正文，介绍客服知识库背景。\n" * 65
            ) + "\n文末特殊规则：海蓝会员退货期限是签收后十三天。"
            word = WordDocument()
            word.add_heading("保修条款", level=1)
            word.add_paragraph("蓝鲸设备保修期限为二十三个月。")
            table = word.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "服务"
            table.cell(0, 1).text = "上门维修"
            buf = io.BytesIO()
            word.save(buf)
            fixtures = [
                ("original.txt", txt.encode(), "text/plain"),
                (
                    "policy.md",
                    b"# Return policy\nReturns are allowed within 17 days.\n## Final rule\nVIP warranty lasts 37 months.",
                    "text/markdown",
                ),
                ("policy.pdf", pdf_bytes(), "application/pdf"),
                (
                    "policy.docx",
                    buf.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
            ]
            docs = []
            for filename, data, mime in fixtures:
                res = await client.post(
                    base + "/documents/upload", files={"file": (filename, data, mime)}
                )
                assert res.status_code == 202, res.text
                body = res.json()
                docs.append(body)
                await worker.run_once()
                res = await client.get(base + "/jobs/" + body["job_id"])
                status = res.json()
                assert status["status"] == "ready", status
                async with get_db_context() as db:
                    rev = await db.get(DocumentRevision, body["revision_id"])
                    collections.add(rev.collection)
                    chunks = (
                        await db.scalars(
                            select(DocumentChunk).where(DocumentChunk.revision_id == rev.id)
                        )
                    ).all()
                    assert all(
                        rev.full_text[c.start_offset : c.end_offset] == c.content for c in chunks
                    )
                    assert chunks
                    if filename.endswith(".pdf"):
                        assert chunks[0].page == 1
                    if filename.endswith(".docx"):
                        assert chunks[0].heading_path == ["保修条款"]
                original = await client.get(base + f"/documents/{body['document_id']}/original")
                assert hashlib.sha256(original.content).digest() == hashlib.sha256(data).digest()
                print("INGEST_OK", filename, len(chunks), flush=True)
            paste = await client.post(
                base + "/documents/text",
                json={"title": "粘贴原文", "text": "紫竹会员取消预约须提前四十八小时。"},
            )
            assert paste.status_code == 202, paste.text
            await worker.run_once()
            paste_job = await client.get(base + "/jobs/" + paste.json()["job_id"])
            assert paste_job.json()["status"] == "ready", paste_job.text
            duplicate = await client.post(base + "/documents/upload", files={"file": fixtures[0]})
            assert duplicate.json()["document_id"] == docs[0]["document_id"], duplicate.text
            hits = await service.search(config_id, "海蓝会员退货期限是多少天？")
            assert any("十三天" in h["content"] for h in hits), hits
            assert not await service.search(other_id, "海蓝会员退货期限是多少天？")
            wrong = await client.get(
                f"/api/v1/admin/knowledge-configs/{other_id}/documents/{docs[0]['document_id']}/content"
            )
            assert wrong.status_code == 404, wrong.text
            print("REAL_RETRIEVAL_TAIL_AND_ISOLATION_OK", flush=True)
            async with get_db_context() as db:
                db.add(
                    Agent(
                        id=agent_id,
                        name="RAG acceptance",
                        type="kb",
                        knowledge_config_id=config_id,
                        system_prompt="你是知识库助手。回答前必须调用 kb_search，准确回答并标注工具给出的 [S数字] 引用。",
                    )
                )
            from app.services.agent.core.service import agent_service
            from app.services.chat_stream_legacy import ChatStreamOrchestrator
            from app.services.chat_stream_sdk import ChatStreamOrchestratorSDK
            from app.services.conversation import ConversationService

            citation_url = None
            for stream_class in [ChatStreamOrchestratorSDK, ChatStreamOrchestrator]:
                conv, user_id, msg_id = uid(), uid(), uid()
                async with get_db_context() as db:
                    db.add(User(id=user_id))
                    await db.flush()
                    db.add(Conversation(id=conv, user_id=user_id))
                    await db.flush()
                async with get_db_context() as db:
                    orchestrator = stream_class(
                        conversation_service=ConversationService(db),
                        agent_service=agent_service,
                        conversation_id=conv,
                        user_id=user_id,
                        user_message="请查询海蓝会员退货期限是多少天？必须检索并引用来源。",
                        user_message_id=uid(),
                        assistant_message_id=msg_id,
                        agent_id=agent_id,
                    )
                    events = []
                    async for event in orchestrator.run():
                        event = event.model_dump() if hasattr(event, "model_dump") else event
                        events.append(event)
                errors = [e for e in events if e["type"] == "error"]
                assert not errors, errors
                cites = [e for e in events if e["type"] == "assistant.citations"]
                assert cites, [e["type"] for e in events]
                final = [e for e in events if e["type"] == "assistant.final"][-1]
                assert (
                    "十三天" in final["payload"]["content"] or "13" in final["payload"]["content"]
                ), final["payload"]["content"]
                assert next(
                    i for i, e in enumerate(events) if e["type"] == "assistant.citations"
                ) < next(i for i, e in enumerate(events) if e["type"] == "assistant.final")
                citation_url = cites[0]["payload"]["items"][0]["source_url"]
                res = await client.get(citation_url)
                assert res.status_code == 200, res.text
                assert res.json()["text"][res.json()["start_offset"] : res.json()["end_offset"]]
                async with get_db_context() as db:
                    msg = await db.get(Message, msg_id)
                    assert msg.extra_metadata["citations"]
                    assert "source_url" not in msg.extra_metadata["citations"][0]
                    meta = await history_metadata(db, msg)
                    assert meta["citations"][0]["source_url"]
                print("REAL_AGENT_STREAM_HISTORY_OK", stream_class.__name__, flush=True)
            res = await client.get(citation_url + "tampered")
            assert res.status_code == 401, res.text
            # Expiry, wrong citation identity and wrong revision ownership.
            from app.models.knowledge_document import MessageCitation
            from app.services.knowledge.rag.security import source_url as sign_url

            async with get_db_context() as db:
                stored = await db.scalar(
                    select(MessageCitation).where(MessageCitation.message_id == msg_id).limit(1)
                )
                old_ttl = settings.KNOWLEDGE_CITATION_TTL
                settings.KNOWLEDGE_CITATION_TTL = -1
                expired = sign_url(stored)
                settings.KNOWLEDGE_CITATION_TTL = old_ttl
            assert (await client.get(expired)).status_code == 401
            foreign_revision = await client.get(
                base + f"/documents/{docs[0]['document_id']}/content",
                params={"revision_id": docs[1]["revision_id"]},
            )
            assert foreign_revision.status_code == 404, foreign_revision.text
            print("EXPIRY_AND_REVISION_OWNERSHIP_OK", flush=True)
            old = docs[0]
            broken = await client.post(
                base + f"/documents/{old['document_id']}/revisions",
                files={"file": ("broken.txt", b"\xff", "text/plain")},
            )
            assert broken.status_code == 202, broken.text
            await worker.run_once()
            assert (await client.get(base + "/jobs/" + broken.json()["job_id"])).json()[
                "status"
            ] == "failed"
            async with get_db_context() as db:
                assert (await db.get(Document, old["document_id"])).current_revision_id == old[
                    "revision_id"
                ]
            print("FAILED_REBUILD_PRESERVES_OLD_VERSION_OK", flush=True)
            replacement = await client.post(
                base + f"/documents/{old['document_id']}/revisions",
                data={"text": "海蓝会员新规则：退货期限改为签收后十五天。"},
            )
            assert replacement.status_code == 202, replacement.text
            # 新内容已构建但旧向量清理失败时，任务不能提前显示 ready。
            real_remove_revision = worker.remove_revision

            async def cleanup_unavailable(*args):
                """只注入旧索引删除故障，不停止共享 Qdrant。"""
                raise ConnectionError("injected old-index cleanup failure")

            worker.remove_revision = cleanup_unavailable
            try:
                await worker.run_once()
            finally:
                worker.remove_revision = real_remove_revision
            async with get_db_context() as db:
                doc = await db.get(Document, old["document_id"])
                assert doc.current_revision_id == replacement.json()["revision_id"]
                job = await db.get(IngestionJob, replacement.json()["job_id"])
                assert job.status == "retry_wait"
                assert (await db.get(DocumentRevision, doc.current_revision_id)).status == "ready"
                await db.execute(
                    update(IngestionJob).where(IngestionJob.id == job.id).values(available_at=0)
                )
            # 恢复仅清理旧向量，禁止再次调用模型、改写正在服务的分块。
            real_model = worker.get_embedding_model

            def forbid_reembedding():
                """重试清理不应重建已发布内容。"""
                raise AssertionError("published content must not be re-embedded")

            worker.get_embedding_model = forbid_reembedding
            try:
                await worker.run_once()
            finally:
                worker.get_embedding_model = real_model
            assert (await client.get(base + "/jobs/" + replacement.json()["job_id"])).json()[
                "status"
            ] == "ready"
            async with qdrant() as q:
                for col in collections:
                    remaining = await q.count(
                        col,
                        count_filter=scope_filter(config_id, revision_id=old["revision_id"]),
                        exact=True,
                    )
                    assert remaining.count == 0
            print("OLD_VECTORS_PHYSICALLY_REMOVED_BEFORE_SUCCESS_OK", flush=True)
            latest_source = await client.get(citation_url)
            assert latest_source.status_code == 200
            latest_body = latest_source.json()
            assert latest_body["revision_id"] == replacement.json()["revision_id"]
            assert latest_body["source_updated"] is True
            assert latest_body["start_offset"] is None and latest_body["excerpt"] is None
            assert "十五天" in latest_body["text"]
            assert "十五天" in (await client.get(latest_body["original_url"])).content.decode()
            async with get_db_context() as db:
                saved = await db.get(Message, msg_id)
                metadata = await history_metadata(db, saved)
                assert metadata["citations"][0]["source_updated"] is True
                assert metadata["citations"][0]["start_offset"] is None
            print("HISTORICAL_CITATION_READS_LATEST_SOURCE_OK", flush=True)
            newhits = await service.search(config_id, "海蓝会员退货期限")
            assert all(h["revision_id"] != old["revision_id"] for h in newhits)
            assert any("十五天" in h["content"] for h in newhits)
            await client.patch(base + f"/documents/{old['document_id']}", json={"enabled": False})
            assert all(
                h["document_id"] != old["document_id"]
                for h in await service.search(config_id, "海蓝会员退货期限")
            )
            assert (await client.get(citation_url)).status_code == 200
            await client.delete(base + f"/documents/{old['document_id']}")
            assert (await client.get(citation_url)).status_code == 410
            await worker.run_once()
            print("VERSION_HISTORY_DISABLE_DELETE_OK", flush=True)
            # Invalid UTF-8 is a permanent parse failure, not a half-published document.
            bad = await client.post(
                base + "/documents/upload",
                files={"file": ("invalid.txt", b"\xff\xfe\x00", "text/plain")},
            )
            assert bad.status_code == 202, bad.text
            await worker.run_once()
            status = (await client.get(base + "/jobs/" + bad.json()["job_id"])).json()
            assert status["status"] == "failed", status
            async with get_db_context() as db:
                d = await db.get(Document, bad.json()["document_id"])
                assert d.current_revision_id is None
            print("PERMANENT_FAILURE_NOT_PUBLISHED_OK", flush=True)
            # Controlled network failure uses actual persisted jobs and old published data.
            retry = await client.post(
                base + "/documents/text",
                json={"title": "恢复验收", "text": "恢复验收规定退款处理时间为九个工作日。"},
            )
            assert retry.status_code == 202, retry.text
            claimed = await worker.claim()
            assert claimed.id == retry.json()["job_id"]
            assert await worker.claim() is None
            async with get_db_context() as db:
                await db.execute(
                    update(IngestionJob).where(IngestionJob.id == claimed.id).values(lease_until=0)
                )
            resumed = await worker.claim()
            assert resumed.id == claimed.id and resumed.token != claimed.token
            try:
                async with get_db_context() as db:
                    await worker.guard(db, claimed.id, claimed.token)
            except worker.LeaseLost:
                pass
            else:
                raise AssertionError("stale worker published")
            await worker.failure(resumed, ConnectionError("injected network outage"))
            async with get_db_context() as db:
                retry_row = await db.get(IngestionJob, resumed.id)
                assert (
                    retry_row.status == "retry_wait"
                    and retry_row.attempts == 1
                    and retry_row.available_at > 0
                )
                await db.execute(
                    update(IngestionJob).where(IngestionJob.id == resumed.id).values(available_at=0)
                )
            await worker.run_once()
            assert (await client.get(base + "/jobs/" + resumed.id)).json()["status"] == "ready"
            print("LEASE_FENCING_AND_TRANSIENT_RECOVERY_OK", flush=True)
            # Manual retry and Qdrant outage are injected without stopping shared services.
            retry_bad = await client.post(base + "/jobs/" + bad.json()["job_id"] + "/retry")
            assert retry_bad.status_code == 202, retry_bad.text
            await worker.run_once()
            assert (await client.get(base + "/jobs/" + bad.json()["job_id"])).json()[
                "status"
            ] == "failed"
            from contextlib import asynccontextmanager

            real_qdrant = worker.qdrant

            @asynccontextmanager
            async def unavailable_qdrant():
                """不停止共享服务，通过适配器注入暂时不可用故障。"""
                raise ConnectionError("injected Qdrant outage")
                yield

            outage = await client.post(
                base + "/documents/text",
                json={"title": "Qdrant恢复", "text": "停机恢复验收：设备交付时间为周三。"},
            )
            worker.qdrant = unavailable_qdrant
            try:
                await worker.run_once()
            finally:
                worker.qdrant = real_qdrant
            status = (await client.get(base + "/jobs/" + outage.json()["job_id"])).json()
            assert status["status"] == "retry_wait", status
            async with get_db_context() as db:
                assert (
                    await db.get(Document, outage.json()["document_id"])
                ).current_revision_id is None
                await db.execute(
                    update(IngestionJob)
                    .where(IngestionJob.id == outage.json()["job_id"])
                    .values(available_at=0)
                )
            await worker.run_once()
            assert (await client.get(base + "/jobs/" + outage.json()["job_id"])).json()[
                "status"
            ] == "ready"
            print("MANUAL_RETRY_AND_QDRANT_OUTAGE_RECOVERY_OK", flush=True)
            # 已托管知识源不能经旧管理接口绕过密钥或直接换集合。
            config_path = f"/api/v1/admin/knowledge-configs/{config_id}"
            assert (
                await client.patch(
                    config_path,
                    json={"name": "unauthorized"},
                    headers={"X-Knowledge-Admin-Key": "wrong"},
                )
            ).status_code == 401
            assert (
                await client.patch(config_path, json={"collection_name": "products"})
            ).status_code == 409
            assert (await client.delete(config_path)).status_code == 409
            settings.KNOWLEDGE_RAG_ENABLED = False
            assert (await client.get(base + "/documents")).status_code == 503
            settings.KNOWLEDGE_RAG_ENABLED = True
            # 显式重建可确认当前全局模型，而不是要求修改受保护的 collection 配置。
            async with get_db_context() as db:
                config = await db.get(KnowledgeConfig, config_id)
                config.embedding_model = "old-model-declaration"
            rebuilt = await client.post(base + f"/documents/{retry.json()['document_id']}/reindex")
            assert rebuilt.status_code == 202, rebuilt.text
            await worker.run_once()
            assert (await client.get(base + "/jobs/" + rebuilt.json()["job_id"])).json()[
                "status"
            ] == "ready"
            await worker.run_once()  # old revision prune
            async with get_db_context() as db:
                assert (
                    await db.get(KnowledgeConfig, config_id)
                ).embedding_model == settings.EMBEDDING_MODEL
            print("LEGACY_CONFIG_GUARDS_ROLLBACK_SWITCH_AND_REINDEX_OK", flush=True)

        finally:
            from app.services.agent.core.service import agent_service

            await agent_service.close()
            from app.services.memory import fact_memory, store

            if store._store_instance is not None:
                await store._store_instance.close()
            if fact_memory._fact_service is not None:
                await fact_memory._fact_service.close()
            # Delete only QA points scoped to our random config IDs; never drop shared collection.
            async with get_db_context() as db:
                collections.update((await db.scalars(select(DocumentRevision.collection))).all())
            async with qdrant() as q:
                from qdrant_client import models

                for col in collections:
                    if await q.collection_exists(col):
                        for cid in [config_id, other_id]:
                            await q.delete(
                                col, models.FilterSelector(filter=scope_filter(cid)), wait=True
                            )
            client_minio = storage._client()
            if client_minio.bucket_exists(settings.KNOWLEDGE_BUCKET):
                for key, _ in await storage.objects():
                    await storage.delete(key)
                await asyncio.to_thread(client_minio.remove_bucket, settings.KNOWLEDGE_BUCKET)
            await provider._provider.close()
            tmp.cleanup()
    print("ALL_LIVE_ACCEPTANCE_OK", flush=True)


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--live",
        action="store_true",
        required=True,
        help="Call configured Embedding/chat providers; incurs normal API usage",
    )
    parser.parse_args()
    try:
        asyncio.run(main())
    except BaseException as exc:
        import traceback

        traceback.print_exception(exc)
        raise SystemExit(1)
